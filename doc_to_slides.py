"""
Multi-Document to Presentation Generator
-----------------------------------------------------------
Converts a primary document into PowerPoint slides (.pptx),
integrating up to multiple reference documents for context/evidence
and adhering strictly to a guideline write-up.

Supports: PDF (.pdf), Word (.docx), Markdown (.md), and Text (.txt)
Works both offline (rule-based) and with AI (Gemini API).
Includes an interactive File Picker UI if run without arguments.
"""

import os
import re
import sys
import json
import argparse
from pathlib import Path
from typing import List, Dict, Optional

# Ensure Windows PowerShell handles UTF-8 output cleanly without crashing
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# PowerPoint generation
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Document extractors
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    from google import genai
    from google.genai import types
except ImportError:
    genai = None


# --- COLOR & DESIGN THEMES ---
THEMES = {
    "slate": {  # Modern dark mode
        "bg_color": RGBColor(15, 23, 42),       # Slate 900 (#0F172A)
        "card_color": RGBColor(30, 41, 59),     # Slate 800 (#1E293B)
        "card_border": RGBColor(51, 65, 85),    # Slate 700 (#334155)
        "accent": RGBColor(99, 102, 241),       # Indigo 500 (#6366F1)
        "accent_light": RGBColor(129, 140, 248),# Indigo 400 (#818CF8)
        "title_color": RGBColor(248, 250, 252), # Off-white (#F8FAFC)
        "body_color": RGBColor(203, 213, 225),  # Slate 300 (#CBD5E1)
        "meta_color": RGBColor(148, 163, 184),  # Slate 400 (#94A3B8)
        "font_title": "Segoe UI",
        "font_body": "Segoe UI",
    },
    "corporate": {  # Clean executive light mode
        "bg_color": RGBColor(248, 250, 252),    # Off-white (#F8FAFC)
        "card_color": RGBColor(255, 255, 255),  # Pure white (#FFFFFF)
        "card_border": RGBColor(226, 232, 240), # Slate 200 (#E2E8F0)
        "accent": RGBColor(14, 116, 144),       # Cyan 700 (#0E7490)
        "accent_light": RGBColor(6, 182, 212),  # Cyan 500 (#06B6D4)
        "title_color": RGBColor(15, 23, 42),    # Slate 900 (#0F172A)
        "body_color": RGBColor(51, 65, 85),     # Slate 700 (#334155)
        "meta_color": RGBColor(100, 116, 139),  # Slate 500 (#64748B)
        "font_title": "Calibri",
        "font_body": "Calibri",
    }
}


def pick_files_with_gui():
    """Opens standard Windows file dialog windows to pick documents without typing paths."""
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.withdraw()  # Hide root window
    root.attributes("-topmost", True)  # Bring dialog to front

    filetypes = [
        ("Supported Documents", "*.pdf *.docx *.txt *.md"),
        ("PDF Documents", "*.pdf"),
        ("Word Documents", "*.docx"),
        ("Text / Markdown", "*.txt *.md"),
        ("All Files", "*.*")
    ]

    print("\n[+] Opening file selector for PRIMARY document...")
    main_file = filedialog.askopenfilename(
        title="1 of 4: Select PRIMARY Document to Convert",
        filetypes=filetypes
    )
    if not main_file:
        print("[!] No main file selected. Exiting.")
        sys.exit(0)

    print(f"    Selected Main File: {main_file}")

    # Reference file 1
    ref1 = filedialog.askopenfilename(
        title="2 of 4: Select REFERENCE Document #1 (or Cancel to skip)",
        filetypes=filetypes
    )
    references = []
    if ref1:
        references.append(ref1)
        print(f"    Selected Reference 1: {ref1}")
        # Reference file 2
        ref2 = filedialog.askopenfilename(
            title="3 of 4: Select REFERENCE Document #2 (or Cancel to skip)",
            filetypes=filetypes
        )
        if ref2:
            references.append(ref2)
            print(f"    Selected Reference 2: {ref2}")

    # Guidelines file
    guidelines_file = filedialog.askopenfilename(
        title="4 of 4: Select GUIDELINES File (or Cancel to use defaults)",
        filetypes=[("Text / Markdown Files", "*.txt *.md"), ("All Files", "*.*")]
    )
    if guidelines_file:
        print(f"    Selected Guidelines: {guidelines_file}")

    # Output presentation location
    output_file = filedialog.asksaveasfilename(
        title="Save Presentation As...",
        defaultextension=".pptx",
        filetypes=[("PowerPoint Presentation", "*.pptx")],
        initialfile="my_presentation.pptx"
    )
    if not output_file:
        output_file = "my_presentation.pptx"

    root.destroy()
    return main_file, references, guidelines_file, output_file


def extract_text_from_file(file_path: Path) -> str:
    """Extracts raw text from PDF, DOCX, MD, or TXT."""
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = file_path.suffix.lower()

    if ext == ".pdf":
        if not PdfReader:
            raise ImportError("Please install pypdf: py -m pip install pypdf")
        reader = PdfReader(str(file_path))
        pages_text = []
        for page in reader.pages:
            txt = page.extract_text() or ""
            if txt.strip():
                pages_text.append(txt.strip())
        return "\n\n".join(pages_text)

    elif ext in [".docx", ".doc"]:
        if not docx:
            raise ImportError("Please install python-docx: py -m pip install python-docx")
        doc = docx.Document(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    elif ext in [".txt", ".md", ".markdown"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()

    else:
        raise ValueError(f"Unsupported file format: {ext}. Supported formats: .pdf, .docx, .txt, .md")


def local_multi_doc_slides(
    main_text: str,
    main_name: str,
    ref_texts: List[Dict[str, str]],
    guidelines_text: str
) -> dict:
    """Offline fallback parser."""
    presentation_title = main_name.replace("_", " ").replace("-", " ").title()
    presentation_subtitle = "Document Synthesis Presentation"

    slides = []

    # Slide 1: Objectives & Guidelines (if guidelines are present)
    if guidelines_text:
        guideline_lines = [
            line.strip().lstrip("-*#• ")
            for line in guidelines_text.splitlines()
            if line.strip() and not line.strip().startswith("#")
        ][:4]
        if guideline_lines:
            slides.append({
                "title": "Presentation Guidelines & Scope",
                "badge": "FRAMEWORK",
                "bullets": guideline_lines,
                "notes": "Presentation objectives and guidelines configured for this deck."
            })

    # Main Content Slides: Segment by paragraphs/sections
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", main_text) if p.strip()]
    chunk_size = max(2, len(paragraphs) // 6) if paragraphs else 2

    for i in range(0, len(paragraphs), chunk_size):
        chunk = paragraphs[i:i + chunk_size]
        first_line = chunk[0].split("\n")[0].strip()

        if len(first_line) < 70 and not first_line.endswith("."):
            slide_title = first_line.lstrip("#*• ").strip()
            chunk[0] = chunk[0][len(first_line):].strip()
        else:
            slide_title = f"{presentation_title} - Part {len(slides) + 1}"

        points = []
        for p in chunk:
            sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", p) if s.strip()]
            for s in sentences:
                clean_s = s.replace("\n", " ").strip()
                if len(clean_s) > 15:
                    points.append(clean_s[:160] + ("..." if len(clean_s) > 160 else ""))
                if len(points) >= 4:
                    break
            if len(points) >= 4:
                break

        if points:
            slides.append({
                "title": slide_title,
                "badge": "CORE ANALYSIS",
                "bullets": points,
                "notes": f"Derived from main document: {main_name}"
            })

        if len(slides) >= 8:
            break

    # Reference Slides: Add highlights from reference files
    for ref in ref_texts:
        ref_name = ref["name"]
        ref_content = ref["text"]
        ref_paras = [p.strip() for p in re.split(r"\n\s*\n", ref_content) if p.strip()]

        ref_points = []
        for p in ref_paras[:4]:
            first_sent = re.split(r"(?<=[.!?])\s+", p)[0].strip()
            if len(first_sent) > 15:
                ref_points.append(first_sent[:150] + ("..." if len(first_sent) > 150 else ""))

        if ref_points:
            slides.append({
                "title": f"Reference Insights: {ref_name}",
                "badge": "SUPPORTING DATA",
                "bullets": ref_points,
                "notes": f"Supporting evidence and comparative data from {ref_name}."
            })

    # Conclusion Slide
    slides.append({
        "title": "Conclusions & Next Steps",
        "badge": "SUMMARY",
        "bullets": [
            f"Synthesized core findings from {main_name}.",
            f"Cross-referenced with {len(ref_texts)} supporting sources.",
            "Aligned with defined project presentation guidelines.",
            "Open for questions and strategic discussions."
        ],
        "notes": "Wrap-up talking points and next action items."
    })

    return {
        "presentation_title": presentation_title,
        "presentation_subtitle": f"Based on {main_name} with {len(ref_texts)} Reference Sources",
        "slides": slides
    }


def ai_multi_doc_slides(
    main_text: str,
    main_name: str,
    ref_texts: List[Dict[str, str]],
    guidelines_text: str,
    api_key: Optional[str] = None
) -> dict:
    """Uses Google Gemini API to synthesize documents and adhere to guidelines."""
    if not genai:
        print("[!] google-genai is not installed. Falling back to local offline mode.")
        return local_multi_doc_slides(main_text, main_name, ref_texts, guidelines_text)

    api_key = api_key or os.getenv("GEMINI_API_KEY")
    if not api_key:
        print("[!] No GEMINI_API_KEY found. Falling back to local offline mode.")
        return local_multi_doc_slides(main_text, main_name, ref_texts, guidelines_text)

    client = genai.Client(api_key=api_key)

    ref_payload = ""
    for idx, ref in enumerate(ref_texts, start=1):
        ref_payload += f"\n--- REFERENCE DOCUMENT {idx} ({ref['name']}) ---\n"
        ref_payload += ref["text"][:6000] + "\n"

    prompt = f"""
You are an expert presentation designer and executive communication strategist.
Your task is to transform a Primary Document into a compelling, professional presentation slide deck.

IMPORTANT INPUTS:
1. GUIDELINES & SPECIFICATIONS:
{guidelines_text if guidelines_text.strip() else "Create a clear, executive-level 6-10 slide presentation with sharp takeaways."}

2. PRIMARY DOCUMENT TO CONVERT ({main_name}):
{main_text[:14000]}

3. SUPPORTING REFERENCE DOCUMENTS (Use for validation, comparisons, context, or evidence):
{ref_payload if ref_payload.strip() else "None provided."}

INSTRUCTIONS:
- Strictly adhere to the GUIDELINES provided (tone, structure, audience, formatting instructions).
- Convert the PRIMARY DOCUMENT's key topics into cohesive slides.
- Weave in insights, data points, or comparisons from the REFERENCE DOCUMENTS where applicable.
- For each slide, write 3-4 concise, high-impact bullet points (avoid walls of text).
- Include helpful presenter talking points in the 'notes' field.
- Add a short 1-2 word category badge for each slide (e.g., "OVERVIEW", "ANALYSIS", "DATA", "STRATEGY", "SUMMARY").

Return ONLY a valid JSON object matching this schema:
{{
  "presentation_title": "Clear, engaging presentation title",
  "presentation_subtitle": "Descriptive subtitle mentioning context or date",
  "slides": [
    {{
      "title": "Slide Title",
      "badge": "CATEGORY",
      "bullets": [
        "First punchy takeaway point",
        "Second point with supporting context or data",
        "Third actionable or strategic insight"
      ],
      "notes": "Speaker notes and talking points for the presenter."
    }}
  ]
}}
"""

    try:
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                temperature=0.3
            )
        )
        return json.loads(response.text)
    except Exception as e:
        print(f"[!] AI generation failed or returned invalid JSON ({e}). Falling back to local parser.")
        return local_multi_doc_slides(main_text, main_name, ref_texts, guidelines_text)


def create_presentation(data: dict, output_path: Path, theme_name: str = "slate"):
    """Creates a modern 16:9 widescreen PowerPoint deck with cards and badges."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    theme = THEMES.get(theme_name, THEMES["slate"])

    # --- TITLE SLIDE ---
    title_slide = prs.slides.add_slide(blank_layout)
    title_slide.background.fill.solid()
    title_slide.background.fill.fore_color.rgb = theme["bg_color"]

    top_bar = title_slide.shapes.add_shape(
        1,  # RECTANGLE
        Inches(1.2), Inches(1.8), Inches(1.5), Inches(0.08)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = theme["accent"]
    top_bar.line.fill.background()

    title_box = title_slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(3.5))
    tf = title_box.text_frame
    tf.word_wrap = True

    p_title = tf.paragraphs[0]
    p_title.text = data.get("presentation_title", "Document Presentation")
    p_title.font.name = theme["font_title"]
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = theme["title_color"]

    p_sub = tf.add_paragraph()
    p_sub.text = data.get("presentation_subtitle", "Multi-Source Document Synthesis")
    p_sub.font.name = theme["font_body"]
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = theme["accent_light"]
    p_sub.space_before = Pt(18)

    # --- CONTENT SLIDES ---
    for idx, slide_data in enumerate(data.get("slides", []), start=1):
        slide = prs.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = theme["bg_color"]

        # Category Badge
        badge_text = slide_data.get("badge", "SECTION").upper()
        badge_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.7), Inches(3.0), Inches(0.4))
        btf = badge_box.text_frame
        btf.word_wrap = False
        bp = btf.paragraphs[0]
        bp.text = f"■  {badge_text}"
        bp.font.name = theme["font_body"]
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = theme["accent"]

        # Slide Title
        header_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.3), Inches(0.9))
        htf = header_box.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = slide_data.get("title", f"Slide {idx}")
        hp.font.name = theme["font_title"]
        hp.font.size = Pt(30)
        hp.font.bold = True
        hp.font.color.rgb = theme["title_color"]

        # Content Card Background
        card = slide.shapes.add_shape(
            1,  # RECTANGLE
            Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = theme["card_color"]
        card.line.color.rgb = theme["card_border"]
        card.line.width = Pt(1)

        # Bullet Content Box
        content_box = slide.shapes.add_textbox(Inches(1.4), Inches(2.5), Inches(10.5), Inches(3.8))
        ctf = content_box.text_frame
        ctf.word_wrap = True

        bullets = slide_data.get("bullets", [])
        for b_idx, bullet in enumerate(bullets):
            para = ctf.paragraphs[0] if b_idx == 0 else ctf.add_paragraph()
            para.text = f"•   {bullet}"
            para.font.name = theme["font_body"]
            para.font.size = Pt(18)
            para.font.color.rgb = theme["body_color"]
            para.space_before = Pt(14)
            para.space_after = Pt(4)

        # Footer / Slide Number
        footer_box = slide.shapes.add_textbox(Inches(11.0), Inches(6.8), Inches(1.5), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.alignment = PP_ALIGN.RIGHT
        fp.text = f"{idx}"
        fp.font.name = theme["font_body"]
        fp.font.size = Pt(12)
        fp.font.color.rgb = theme["meta_color"]

        # Speaker notes
        if slide_data.get("notes"):
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_data["notes"]

    prs.save(str(output_path))
    print(f"\n[OK] Presentation created successfully!")
    print(f"     File: {output_path.resolve()}")
    print(f"     Slides: {len(data.get('slides', [])) + 1} (including title)")


def main():
    parser = argparse.ArgumentParser(
        description="Extract a primary document, cross-reference 2 supporting files, apply guidelines, and create slides."
    )
    parser.add_argument("main_file", nargs="?", default=None, help="Primary file to convert (.pdf, .docx, .txt, .md)")
    parser.add_argument(
        "-r", "--references", nargs="+", default=[],
        help="One or two reference documents (.pdf, .docx, .txt, .md)"
    )
    parser.add_argument(
        "-g", "--guidelines", type=str, default=None,
        help="Path to guideline file (.txt, .md)"
    )
    parser.add_argument(
        "--guideline-text", type=str, default="",
        help="Direct inline guideline instructions as a string"
    )
    parser.add_argument(
        "-o", "--output", type=str, default=None,
        help="Output PowerPoint path (default: presentation.pptx)"
    )
    parser.add_argument(
        "--theme", choices=["slate", "corporate"], default="slate",
        help="Visual theme: 'slate' (dark modern) or 'corporate' (clean light)"
    )
    parser.add_argument(
        "--use-ai", action="store_true",
        help="Use Gemini AI to synthesize documents and apply guidelines"
    )
    parser.add_argument(
        "--api-key", type=str, default=None,
        help="Gemini API Key (optional, defaults to GEMINI_API_KEY env var)"
    )

    args = parser.parse_args()

    # If no main_file is provided on the command line, launch standard Windows File Picker GUI
    if not args.main_file:
        print("[*] No file paths specified. Launching interactive File Picker...")
        main_path_str, ref_paths_list, guidelines_path_str, output_path_str = pick_files_with_gui()
        main_path = Path(main_path_str)
        ref_args = ref_paths_list
        guidelines_path = Path(guidelines_path_str) if guidelines_path_str else None
        output_path = Path(output_path_str)
    else:
        main_path = Path(args.main_file)
        ref_args = args.references
        guidelines_path = Path(args.guidelines) if args.guidelines else None
        output_path = Path(args.output or "presentation.pptx")

    # 1. Read Primary File
    print(f"\n[*] Extracting primary document: {main_path.name}...")
    main_text = extract_text_from_file(main_path)
    print(f"    -> Extracted {len(main_text)} characters.")

    # 2. Read Reference Files
    ref_texts = []
    for ref_arg in ref_args:
        ref_path = Path(ref_arg)
        print(f"[*] Extracting reference document: {ref_path.name}...")
        text = extract_text_from_file(ref_path)
        ref_texts.append({"name": ref_path.name, "text": text})
        print(f"    -> Extracted {len(text)} characters from {ref_path.name}.")

    # 3. Read Guidelines
    guidelines = args.guideline_text
    if guidelines_path:
        print(f"[*] Loading guidelines from: {guidelines_path.name}...")
        g_file_text = extract_text_from_file(guidelines_path)
        guidelines = f"{guidelines}\n{g_file_text}".strip()

    if guidelines:
        print(f"[*] Guidelines active ({len(guidelines)} characters).")
    else:
        print("[*] No custom guidelines specified. Using standard executive presentation format.")

    # 4. Generate Deck Outline
    if args.use_ai:
        print("\n[*] Synthesizing presentation with Gemini AI...")
        data = ai_multi_doc_slides(main_text, main_path.stem, ref_texts, guidelines, args.api_key)
    else:
        print("\n[*] Synthesizing presentation using local multi-document parser...")
        data = local_multi_doc_slides(main_text, main_path.stem, ref_texts, guidelines)

    # 5. Build PowerPoint Deck
    create_presentation(data, output_path, theme_name=args.theme)


if __name__ == "__main__":
    main()
